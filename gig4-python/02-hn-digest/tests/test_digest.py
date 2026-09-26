"""Offline tests: filtering, outputs, email (to a local SMTP server) and the Sheets push (faked).
Run: python -m pytest -q
"""
from __future__ import annotations

import socket
import sys
import threading
import types
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

import hn_digest as hd  # noqa: E402


def story(rank: int, title: str, score: int, domain: str = "example.com") -> hd.Story:
    return hd.Story(rank=rank, id=1000 + rank, title=title, domain=domain, url=f"https://{domain}/{rank}",
                    score=score, comments=3, posted_utc="2026-09-26 10:00", age_hours=2.0, type="story",
                    hn_url=f"https://news.ycombinator.com/item?id={1000 + rank}", matched="")


STORIES = [
    story(1, "Show HN: A Python tool for PDFs", 150),
    story(2, "What she said about the market", 300),       # "said" must NOT match "ai"
    story(3, "AI agents in production", 90),
    story(4, "Rust 2.0 released", 10),                      # below min score
    story(5, "Why we left the cloud", 60, domain="python.org"),  # matches via domain
]


def test_filter_whole_words_domain_and_min_score() -> None:
    matches = hd.apply_filter([hd.Story(**vars(s)) for s in STORIES], ["python", "ai", "rust"], min_score=20)
    assert [m.rank for m in matches] == [1, 3, 5]  # sorted by score, desc
    assert matches[0].matched == "python"
    assert matches[1].matched == "ai"


def test_outputs_and_summary(tmp_path: Path) -> None:
    stories = [hd.Story(**vars(s)) for s in STORIES]
    matches = hd.apply_filter(stories, ["python", "ai"], 0)
    hd.write_csv(matches, tmp_path / "m.csv")
    hd.write_excel(stories, matches, ["python", "ai"], 0, tmp_path / "r.xlsx")
    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb.sheetnames == ["Matches", "Keyword summary", "All top stories", "Run info"]
    summary = {r[0]: r[1] for r in wb["Keyword summary"].iter_rows(min_row=2, values_only=True)}
    assert summary == {"python": 2, "ai": 1}
    assert (tmp_path / "m.csv").read_text(encoding="utf-8").startswith("rank,id,title")


def test_email_has_text_and_html() -> None:
    matches = hd.apply_filter([hd.Story(**vars(s)) for s in STORIES], ["python"], 0)
    msg = hd.build_email(matches, ["python"], "from@example.com", ["a@example.com", "b@example.com"])
    assert msg["To"] == "a@example.com, b@example.com"
    assert "Show HN: A Python tool for PDFs" in msg.get_body(("plain",)).get_content()
    assert "<table" in msg.get_body(("html",)).get_content()


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def test_send_email_to_local_smtp_server(monkeypatch: pytest.MonkeyPatch) -> None:
    """Spin up a minimal SMTP server and check the message really arrives."""
    received: list[bytes] = []
    port = _free_port()
    srv = socket.socket()
    srv.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    srv.bind(("127.0.0.1", port))
    srv.listen(1)

    def serve() -> None:
        conn, _ = srv.accept()
        f = conn.makefile("rb")
        conn.sendall(b"220 test ESMTP\r\n")
        data_mode, buf = False, []
        for line in f:
            if data_mode:
                if line in (b".\r\n", b".\n"):
                    received.append(b"".join(buf))
                    data_mode = False
                    conn.sendall(b"250 OK queued\r\n")
                else:
                    buf.append(line)
                continue
            cmd = line.strip().upper()
            if cmd.startswith(b"EHLO") or cmd.startswith(b"HELO"):
                conn.sendall(b"250 test\r\n")
            elif cmd == b"DATA":
                data_mode = True
                conn.sendall(b"354 go ahead\r\n")
            elif cmd == b"QUIT":
                conn.sendall(b"221 bye\r\n")
                break
            else:
                conn.sendall(b"250 OK\r\n")
        conn.close()

    thread = threading.Thread(target=serve, daemon=True)
    thread.start()
    monkeypatch.setenv("SMTP_HOST", "127.0.0.1")
    monkeypatch.setenv("SMTP_PORT", str(port))
    monkeypatch.setenv("SMTP_SECURITY", "none")
    monkeypatch.delenv("SMTP_USER", raising=False)
    matches = hd.apply_filter([hd.Story(**vars(s)) for s in STORIES], ["python"], 0)
    hd.send_email(hd.build_email(matches, ["python"], "from@example.com", ["to@example.com"]))
    thread.join(timeout=5)
    srv.close()
    assert received and b"Subject: HN digest" in received[0]


def test_send_email_requires_host(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("SMTP_HOST", raising=False)
    with pytest.raises(hd.ConfigError):
        hd.send_email(hd.build_email([], ["x"], "a@example.com", ["b@example.com"]))


def test_google_sheet_push_with_fake_client(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    calls: dict[str, object] = {}

    class FakeWorksheet:
        def clear(self) -> None:
            calls["cleared"] = True

        def update(self, values: list, range_name: str) -> None:
            calls["values"], calls["range"] = values, range_name

    class NotFound(Exception):
        pass

    class FakeSheet:
        def worksheet(self, name: str) -> FakeWorksheet:
            raise NotFound(name)

        def add_worksheet(self, title: str, rows: int, cols: int) -> FakeWorksheet:
            calls["added"] = (title, cols)
            return FakeWorksheet()

    fake = types.SimpleNamespace(
        WorksheetNotFound=NotFound,
        service_account=lambda filename: types.SimpleNamespace(open_by_key=lambda key: FakeSheet()),
    )
    monkeypatch.setitem(sys.modules, "gspread", fake)
    key = tmp_path / "key.json"
    key.write_text("{}")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", str(key))
    matches = hd.apply_filter([hd.Story(**vars(s)) for s in STORIES], ["python"], 0)
    assert hd.push_to_google_sheet(matches, "sheet-id", "HN matches") == 2
    assert calls["added"] == ("HN matches", len(hd.HEADERS))
    assert calls["values"][0] == hd.HEADERS and calls["range"] == "A1" and calls["cleared"]


def test_bad_arguments_fail_clearly() -> None:
    with pytest.raises(hd.ConfigError):
        hd.resolve(["--top", "900", "--config", "missing.toml"])
    with pytest.raises(hd.ConfigError):
        hd.resolve(["--send", "--config", "missing.toml"])
