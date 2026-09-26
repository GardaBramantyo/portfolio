"""Tests: parsing helpers, organize (dry run / apply / undo / re-run) and merge on the demo inbox.
Run: python -m pytest -q
"""
from __future__ import annotations

import shutil
import sys
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

import file_automator as fa  # noqa: E402
import make_sample_inbox  # noqa: E402


@pytest.mark.parametrize("raw, expected", [
    ("1.990,00", 1990.0), ("$1,250.00", 1250.0), ("596,00", 596.0), ("1,234", 1234.0),
    ("(12.50)", -12.5), (" 49 ", 49.0), (7, 7.0),
])
def test_parse_number(raw: object, expected: float) -> None:
    assert fa.parse_number(raw) == expected


def test_parse_number_rejects_text() -> None:
    with pytest.raises(ValueError):
        fa.parse_number("n/a")


def test_parse_date_formats() -> None:
    fmts = ["%Y-%m-%d", "%d/%m/%Y"]
    assert fa.parse_date("2026-01-05", fmts) == date(2026, 1, 5)
    assert fa.parse_date("15/02/2026", fmts) == date(2026, 2, 15)
    with pytest.raises(ValueError):
        fa.parse_date("31/02/2026", fmts)


@pytest.mark.parametrize("stem, slug", [
    ("Copy of March export", "march-export"), ("Invoice #443 (copy)", "invoice-443"),
    ("sales_feb-2026 FINAL (2)", "sales-feb-2026-final"), ("IMG_2041", "img-2041"), ("???", "file"),
])
def test_slugify(stem: str, slug: str) -> None:
    assert fa.slugify(stem) == slug


@pytest.fixture()
def workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A private copy of the project with a fresh demo inbox."""
    monkeypatch.setattr(make_sample_inbox, "INBOX", tmp_path / "sample_inbox")
    make_sample_inbox.main()
    shutil.copy(ROOT / "config.toml", tmp_path / "config.toml")
    monkeypatch.chdir(tmp_path)
    return tmp_path


def test_dry_run_changes_nothing(workspace: Path) -> None:
    assert fa.main(["organize"]) == 0
    assert not (workspace / "output").exists()


def test_run_organizes_merges_and_is_repeatable(workspace: Path) -> None:
    assert fa.main(["run"]) == 0
    org = workspace / "output" / "organized"
    assert sorted(p.name for p in (org / "Spreadsheets").iterdir()) == [
        "2026-02-02_sales-jan-2026.csv", "2026-03-03_sales-feb-2026-final.csv", "2026-04-01_march-export.xlsx"]
    assert [p.name for p in (org / "Duplicates").iterdir()] == ["Copy of March export.xlsx"]
    assert not any(p.name.startswith("~$") for p in org.rglob("*"))

    wb = load_workbook(workspace / "output" / "merged_sales.xlsx")
    assert wb.sheetnames == ["Summary", "Data", "Rejected", "Duplicates removed"]
    data = list(wb["Data"].iter_rows(min_row=2, values_only=True))
    assert len(data) == 17
    assert len({r[0] for r in data}) == 17                         # no duplicate order ids
    assert {r[3] for r in data} == {"North", "South", "East", "West"}  # regions cleaned
    assert round(sum(r[6] for r in data), 2) == 8889.0
    reasons = [r[-1] for r in wb["Rejected"].iter_rows(min_row=2, values_only=True)]
    assert any("amount is empty" in r for r in reasons) and any("unrecognised date" in r for r in reasons)
    assert "Café Lumière" in {r[2] for r in data}                   # cp1252 file decoded correctly

    # A second scheduled run must not duplicate anything.
    before = sorted(p.relative_to(org).as_posix() for p in org.rglob("*") if "_logs" not in p.parts)
    assert fa.main(["run"]) == 0
    after = sorted(p.relative_to(org).as_posix() for p in org.rglob("*") if "_logs" not in p.parts)
    assert before == after


def test_undo_reverts_copies(workspace: Path) -> None:
    assert fa.main(["organize", "--apply"]) == 0
    log_file = next((workspace / "output" / "organized" / "_logs").glob("organize_*.csv"))
    assert fa.main(["undo", str(log_file)]) == 0
    remaining = [p for p in (workspace / "output" / "organized").rglob("*") if p.is_file() and "_logs" not in p.parts]
    assert remaining == []
    assert len(list((workspace / "sample_inbox").iterdir())) == 10  # originals untouched


def test_move_mode_empties_inbox(workspace: Path) -> None:
    assert fa.main(["organize", "--apply", "--mode", "move"]) == 0
    assert [p.name for p in (workspace / "sample_inbox").iterdir()] == ["~$March export.xlsx"]


def test_missing_config_and_folder_errors(workspace: Path) -> None:
    assert fa.main(["--config", "nope.toml", "merge"]) == 2
    assert fa.main(["merge", "--source", "does-not-exist"]) == 2


def test_run_lock_blocks_overlap(workspace: Path) -> None:
    lock = workspace / "logs" / "run.lock"
    lock.parent.mkdir(parents=True, exist_ok=True)
    lock.write_text("123")
    assert fa.main(["run"]) == 2
