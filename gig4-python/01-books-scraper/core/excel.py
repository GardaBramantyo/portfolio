"""Excel writer helper (openpyxl): formatted tables, summary sheets, sensible widths.

Usage:
    from openpyxl import Workbook
    from core.excel import write_table, write_key_values, save_workbook

    wb = Workbook()
    write_table(wb, "Books", ["Title", "Price"], rows,
                number_formats={"Price": '"£"#,##0.00'})
    write_key_values(wb, "Run info", [("Rows", len(rows)), ("Source", url)])
    save_workbook(wb, "output/books.xlsx")

Every table sheet gets: bold header on a dark fill, frozen header row, filter
buttons (an Excel Table with banded rows), and column widths fitted to content.
No formulas are written, so the file opens identically in Excel, LibreOffice and
Google Sheets.
"""
from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
MAX_WIDTH = 60
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # characters Excel rejects


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = _ILLEGAL.sub("", value)
        # Stop spreadsheet apps from executing cell text as a formula (CSV/formula injection).
        if value[:1] in ("=", "+", "-", "@") and not _looks_numeric(value):
            value = "'" + value
    return value


def _looks_numeric(text: str) -> bool:
    try:
        float(text)
        return True
    except ValueError:
        return False


def _unique_sheet(wb: Workbook, title: str) -> Worksheet:
    """Use the default empty sheet if it's untouched, otherwise add a new one."""
    title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31]
    first = wb.worksheets[0]
    if len(wb.worksheets) == 1 and first.title == "Sheet" and not first._cells:  # untouched default sheet
        first.title = title
        return first
    return wb.create_sheet(title)


def _table_name(title: str, wb: Workbook) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "_", title) or "Table"
    if base[0].isdigit():
        base = "T_" + base
    existing = {t for ws in wb.worksheets for t in ws.tables}
    name, n = base, 2
    while name in existing:
        name, n = f"{base}_{n}", n + 1
    return name


def autosize(ws: Worksheet, min_width: int = 8, max_width: int = MAX_WIDTH) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            value = cell.value
            if isinstance(value, (datetime, date)):
                length = 19 if isinstance(value, datetime) else 10
            elif isinstance(value, float):
                length = len(f"{value:,.2f}")
            else:
                length = max(len(line) for line in str(value).splitlines() or [""])
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 2))


def write_table(
    wb: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    number_formats: dict[str, str] | None = None,
    widths: dict[str, int] | None = None,
    style: str = "TableStyleMedium2",
) -> Worksheet:
    """Write headers + rows as a formatted Excel Table on a new sheet."""
    ws = _unique_sheet(wb, title)
    ws.append(list(headers))
    count = 0
    for row in rows:
        ws.append([_clean(v) for v in row])
        count += 1
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    last_col = get_column_letter(len(headers))
    if count:
        table = Table(displayName=_table_name(title, wb), ref=f"A1:{last_col}{count + 1}")
        table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A1:{last_col}1"

    formats = number_formats or {}
    for idx, header in enumerate(headers, start=1):
        fmt = formats.get(header)
        if fmt is None:
            continue
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.number_format = fmt
    for idx, header in enumerate(headers, start=1):
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            if isinstance(cell.value, datetime) and header not in formats:
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date) and header not in formats:
                cell.number_format = "yyyy-mm-dd"

    autosize(ws)
    for header, width in (widths or {}).items():
        if header in headers:
            ws.column_dimensions[get_column_letter(list(headers).index(header) + 1)].width = width
    return ws


def write_key_values(wb: Workbook, title: str, items: Iterable[tuple[str, Any]], heading: str | None = None) -> Worksheet:
    """A two-column 'label: value' sheet for run info or headline numbers."""
    ws = _unique_sheet(wb, title)
    start = 1
    if heading:
        ws["A1"] = heading
        ws["A1"].font = TITLE_FONT
        start = 3
    for offset, (label, value) in enumerate(items):
        r = start + offset
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=_clean(value))
        cell.alignment = Alignment(horizontal="left")
        if isinstance(value, float):
            cell.number_format = "#,##0.00"
        elif isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"
    autosize(ws)
    return ws


def save_workbook(wb: Workbook, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError as exc:
        raise PermissionError(f"Cannot write {path}. Is it open in Excel? Close it and run again.") from exc
    return path
