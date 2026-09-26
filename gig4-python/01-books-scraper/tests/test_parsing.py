"""Offline tests: parsing and output, no network needed.  Run: python -m pytest -q"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

import books_scraper as bs  # noqa: E402

LISTING = """<html><head><meta charset="utf-8"></head><body>
<div class="side_categories"><ul class="nav nav-list"><li><a href="books_1/index.html">Books</a>
<ul><li><a href="catalogue/category/books/travel_2/index.html">Travel</a></li>
<li><a href="catalogue/category/books/mystery_3/index.html">Mystery</a></li></ul></li></ul></div>
<article class="product_pod"><div class="image_container"><a href="../../../sharp-objects_997/index.html">
<img src="../../../../media/cache/32/51/3251cf3a3412f53f339e42cac2134093.jpg" alt="Sharp Objects"></a></div>
<p class="star-rating Four"></p><h3><a href="../../../sharp-objects_997/index.html" title="Sharp Objects">Sharp Objects</a></h3>
<div class="product_price"><p class="price_color">£47.82</p><p class="instock availability">
<i class="icon-ok"></i> In stock</p></div></article>
<ul class="pager"><li class="current">Page 1 of 2</li><li class="next"><a href="page-2.html">next</a></li></ul>
</body></html>""".encode("utf-8")

DETAIL = """<html><head><meta charset="utf-8"></head><body><table class="table table-striped">
<tr><th>UPC</th><td>e00eb4fd7b871a48</td></tr><tr><th>Product Type</th><td>Books</td></tr>
<tr><th>Price (excl. tax)</th><td>£47.82</td></tr><tr><th>Availability</th><td>In stock (20 available)</td></tr>
<tr><th>Number of reviews</th><td>0</td></tr></table></body></html>""".encode("utf-8")

PAGE_URL = "https://books.toscrape.com/catalogue/category/books/mystery_3/index.html"


def test_parse_categories() -> None:
    cats = bs.parse_categories(LISTING, "https://books.toscrape.com/")
    assert cats == [
        ("Travel", "https://books.toscrape.com/catalogue/category/books/travel_2/index.html"),
        ("Mystery", "https://books.toscrape.com/catalogue/category/books/mystery_3/index.html"),
    ]


def test_parse_listing_and_next_page() -> None:
    books, next_url = bs.parse_listing(LISTING, PAGE_URL, "Mystery")
    assert len(books) == 1
    book = books[0]
    assert book.title == "Sharp Objects"
    assert book.price_gbp == 47.82
    assert book.rating == 4
    assert book.in_stock is True
    assert book.product_url == "https://books.toscrape.com/catalogue/sharp-objects_997/index.html"
    assert next_url == "https://books.toscrape.com/catalogue/category/books/mystery_3/page-2.html"


def test_detail_sets_stock_and_upc() -> None:
    book = bs.parse_listing(LISTING, PAGE_URL, "Mystery")[0][0]
    bs.apply_detail(book, bs.parse_detail(DETAIL))
    assert (book.upc, book.stock_qty, book.reviews) == ("e00eb4fd7b871a48", 20, 0)


def test_parse_price_rejects_garbage() -> None:
    assert bs.parse_price("£1,051.77") == 1051.77
    try:
        bs.parse_price("n/a")
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def test_outputs(tmp_path: Path) -> None:
    books = bs.parse_listing(LISTING, PAGE_URL, "Mystery")[0]
    args = bs.resolve_args(["--out-dir", str(tmp_path)])
    bs.write_csv(books, tmp_path / "books.csv")
    bs.write_excel(books, bs.RunStats(categories=1, listing_pages=1), args, 1.0, tmp_path / "books.xlsx")
    rows = list(csv.DictReader((tmp_path / "books.csv").open(encoding="utf-8-sig")))
    assert rows[0]["title"] == "Sharp Objects"
    wb = load_workbook(tmp_path / "books.xlsx")
    assert wb.sheetnames == ["Books", "By category", "Run info"]
    assert wb["Books"]["A1"].value == "Category"
    assert wb["Books"]["C2"].value == 47.82


def test_delay_floor() -> None:
    try:
        bs.resolve_args(["--delay", "0.1"])
    except bs.ConfigError:
        return
    raise AssertionError("expected ConfigError for a too-small delay")


class FlakyHttp:
    """Stub client: every URL fails on its first request, then succeeds."""

    def __init__(self) -> None:
        self.calls: dict[str, int] = {}

    def get(self, url: str):  # noqa: ANN201
        import types

        self.calls[url] = self.calls.get(url, 0) + 1
        if self.calls[url] == 1:
            raise ConnectionError("simulated network blip")
        return types.SimpleNamespace(content=DETAIL if "sharp-objects" in url else LISTING.replace(b'<li class="next"><a href="page-2.html">next</a></li>', b""))


def test_second_pass_recovers_failed_pages() -> None:
    http = FlakyHttp()
    stats = bs.RunStats()
    args = bs.resolve_args(["--retry-pause", "0"])
    books = list(bs.iter_category(http, "Mystery", PAGE_URL, None, stats))  # listing fails once
    assert books == [] and len(stats.failed_listings) == 1
    seen: set[str] = set()
    bs.second_pass(http, args, stats, books, seen)  # listing recovers; its product page fails once...
    assert len(books) == 1
    assert stats.errors == 1 and books[0].stock_qty is None  # ...and the final attempt of that detail is the 1st call
    book = books[0]
    stats2 = bs.RunStats()
    assert bs.fetch_detail(http, book, stats2) is True and book.stock_qty == 20  # later call succeeds


def test_second_pass_recovers_failed_product_page() -> None:
    http = FlakyHttp()
    stats = bs.RunStats()
    book = bs.parse_listing(LISTING, PAGE_URL, "Mystery")[0][0]
    assert bs.fetch_detail(http, book, stats) is False and stats.failed_details == [book]
    bs.second_pass(http, bs.resolve_args(["--retry-pause", "0"]), stats, [book], set())
    assert (stats.recovered, stats.errors, book.stock_qty) == (1, 0, 20)
